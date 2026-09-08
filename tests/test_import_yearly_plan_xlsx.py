from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "import" / "import_yearly_plan_xlsx.py"
SPEC = importlib.util.spec_from_file_location("import_yearly_plan_xlsx", SCRIPT_PATH)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def create_workbook(path: Path, plan_id: int = 42, grade: str = "9. Sınıf") -> None:
    workbook = Workbook()
    sheet = workbook.active
    metadata = [
        ("Ders Grubu", "Kültür Dersleri"),
        ("Branş/Alan", "Fizik"),
        ("Ders", "Fizik (FL) - 9"),
        ("Sınıf", grade),
        ("Yıl", "2026-2027"),
        ("Sistem", "Örgün"),
        ("Yayınevi", "MEB"),
        ("Kaynak", "https://example.test/plan/42"),
        ("Plan No", str(plan_id)),
    ]
    for row, values in enumerate(metadata, start=1):
        sheet.cell(row, 1, values[0])
        sheet.cell(row, 2, values[1])
    header_row = len(metadata) + 2
    headers = [
        "Hafta",
        "Ders Tarihi",
        "Ders Saati",
        "Ünite/Tema/Öğrenme Alanı",
        "Konu (İçerik Çerçevesi)",
        "Öğrenme Çıktısı Kodları",
        "Öğrenme Çıktısı (Kazanımlar)",
        "Sosyal-Duygusal Becerileri",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(header_row, column, header)
    values = [1, "14-18 Eylül", 2, "Kuvvet", "Hareket", "FİZ.9.1", "Çözüm üretir.", "İletişim"]
    for column, value in enumerate(values, start=1):
        sheet.cell(header_row + 1, column, value)
    workbook.save(path)


class ImportYearlyPlanTests(unittest.TestCase):
    def test_header_mapping_and_extra_fields(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "plan.xlsx"
            create_workbook(path)
            payload, manifest = MODULE.import_workbook(path, root, "2026-2027")

        self.assertEqual(payload["sinif"], "9")
        self.assertEqual(payload["ders"], "Fizik (FL)")
        self.assertEqual(payload["plan"][0]["ÜNİTE"], "Kuvvet")
        self.assertEqual(payload["plan"][0]["KAZANIM"], "Çözüm üretir.")
        self.assertEqual(payload["plan"][0]["SDB"], "İletişim")
        self.assertEqual(payload["plan"][0]["EK_ALANLAR"]["Öğrenme Çıktısı Kodları"], "FİZ.9.1")
        self.assertEqual(manifest["id"], 42)
        self.assertEqual(len(manifest["sha256"]), 64)

    def test_grade_normalization(self) -> None:
        self.assertEqual(MODULE.normalize_grade("12. Sınıf"), "12")
        self.assertEqual(MODULE.normalize_grade("Hazırlık Sınıfı"), "Hazırlık")
        with self.assertRaises(MODULE.ImportValidationError):
            MODULE.normalize_grade("Okul Öncesi")

    def test_duplicate_id_fails_before_publish(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            create_workbook(root / "bir.xlsx", plan_id=7)
            create_workbook(root / "iki.xlsx", plan_id=7)
            old_root = MODULE.REPO_ROOT
            try:
                MODULE.REPO_ROOT = root / "repo"
                MODULE.PUBLIC_DIR = MODULE.REPO_ROOT / "public"
                MODULE.ARCHIVE_DIR = MODULE.REPO_ROOT / "archive"
                with self.assertRaisesRegex(MODULE.ImportValidationError, "Tekrarlı Plan No"):
                    MODULE.run_import(root, "2026-2027", 2, 2)
                self.assertFalse(MODULE.PUBLIC_DIR.exists())
            finally:
                MODULE.REPO_ROOT = old_root
                MODULE.PUBLIC_DIR = old_root / "public"
                MODULE.ARCHIVE_DIR = old_root / "archive"

    def test_bad_workbook_fails(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "bad.xlsx"
            path.write_text("not an xlsx", encoding="utf-8")
            with self.assertRaisesRegex(MODULE.ImportValidationError, "XLSX açılamadı"):
                MODULE.import_workbook(path, root, "2026-2027")

    def test_unchanged_source_is_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source"
            source.mkdir()
            create_workbook(source / "plan.xlsx")
            old_root = MODULE.REPO_ROOT
            try:
                MODULE.REPO_ROOT = root / "repo"
                MODULE.PUBLIC_DIR = MODULE.REPO_ROOT / "public"
                MODULE.ARCHIVE_DIR = MODULE.REPO_ROOT / "archive"
                MODULE.run_import(source, "2026-2027", 1, 1)
                first = (MODULE.PUBLIC_DIR / "years" / "2026-2027" / "index.json").read_bytes()
                MODULE.run_import(source, "2026-2027", 1, 1)
                second = (MODULE.PUBLIC_DIR / "years" / "2026-2027" / "index.json").read_bytes()
                self.assertEqual(first, second)
            finally:
                MODULE.REPO_ROOT = old_root
                MODULE.PUBLIC_DIR = old_root / "public"
                MODULE.ARCHIVE_DIR = old_root / "archive"


if __name__ == "__main__":
    unittest.main()
